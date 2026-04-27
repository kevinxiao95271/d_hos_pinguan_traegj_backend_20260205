"""
导出所有项目的缴费回执照片，并附带报名人姓名、手机号、机构名称。

输出结构：
  export_payment_proofs/
    summary.csv              # 汇总表（报名ID、项目名、机构、报名人、手机、文件名、下载状态）
    files/
      {reg_id}_{机构名}_{报名人}/
        {原始文件名}         # 实际下载的图片/文件

用法：
  pip install pymysql minio openpyxl
  python export_payment_proofs.py

可选参数（直接修改下方 CONFIG）：
  COMPETITION_ID  - 填 None 则导出所有竞赛；填数字则只导出指定竞赛
  OUTPUT_DIR      - 输出目录，默认当前目录下 export_payment_proofs/
  SKIP_DOWNLOAD   - True 则只生成 summary.csv，不下载文件
"""

import sys, os, csv, re
sys.stdout.reconfigure(encoding='utf-8')

import pymysql
from minio import Minio
from minio.error import S3Error

# ─── CONFIG ───────────────────────────────────────────────────────────────────
DB_HOST     = 'gz-cdb-bq7gk3k5.sql.tencentcdb.com'
DB_PORT     = 63606
DB_USER     = 'root'
DB_PASSWORD = 'Yiguo9527_'
DB_NAME     = 'd_hos_pinguan_traegj_20260205'

MINIO_ENDPOINT   = '119.167.165.27:58010'   # 去掉 http://
MINIO_ACCESS_KEY = 'minioadmin'
MINIO_SECRET_KEY = 'Ygcx2025'
MINIO_BUCKET     = 'registration-files'
MINIO_SECURE     = False                     # http → False；https → True

COMPETITION_ID = None        # None = 所有竞赛；数字 = 仅某个竞赛
OUTPUT_DIR     = os.path.join(os.path.dirname(__file__), 'export_payment_proofs')
SKIP_DOWNLOAD  = False       # True 只导出 CSV，不下载图片
# ──────────────────────────────────────────────────────────────────────────────


def safe_name(s: str) -> str:
    """把字符串变成合法的目录/文件名片段"""
    s = str(s).strip()
    return re.sub(r'[\\/:*?"<>|]', '_', s)


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    files_dir = os.path.join(OUTPUT_DIR, 'files')
    os.makedirs(files_dir, exist_ok=True)

    # ── 1. 查数据库 ──────────────────────────────────────────────────────────
    print('连接数据库...')
    conn = pymysql.connect(
        host=DB_HOST, port=DB_PORT,
        user=DB_USER, password=DB_PASSWORD,
        database=DB_NAME, charset='utf8mb4'
    )
    cur = conn.cursor()

    comp_filter = ''
    params = ()
    if COMPETITION_ID is not None:
        comp_filter = 'AND r.competition_id = %s'
        params = (COMPETITION_ID,)

    sql = f"""
        SELECT
            r.id                AS reg_id,
            c.name              AS competition_name,
            r.project_name      AS project_name,
            r.group_type        AS group_type,
            r.status            AS reg_status,
            inst.name           AS institution_name,
            ua.name             AS applicant_name,
            ua.phone            AS applicant_phone,
            mf.id               AS file_id,
            mf.file_name        AS file_name,
            mf.file_url         AS file_url,
            mf.uploaded_at      AS uploaded_at
        FROM material_files mf
        JOIN registrations r    ON mf.registration_id = r.id
        JOIN competitions c     ON r.competition_id   = c.id
        JOIN institutions inst  ON r.institution_id   = inst.id
        JOIN user_accounts ua   ON r.applicant_id     = ua.id
        WHERE mf.type = 'payment_proof'
        {comp_filter}
        ORDER BY r.id, mf.id
    """
    cur.execute(sql, params)
    rows = cur.fetchall()
    cols = [d[0] for d in cur.description]
    conn.close()

    if not rows:
        print('未找到任何缴费回执记录（material_files.type = payment_proof）。')
        return

    print(f'共找到 {len(rows)} 条缴费回执记录，涉及报名数：'
          f'{len(set(r[cols.index("reg_id")] for r in rows))} 个')

    # ── 2. 初始化 MinIO 客户端 ────────────────────────────────────────────────
    minio_client = None
    if not SKIP_DOWNLOAD:
        print('连接 MinIO...')
        minio_client = Minio(
            MINIO_ENDPOINT,
            access_key=MINIO_ACCESS_KEY,
            secret_key=MINIO_SECRET_KEY,
            secure=MINIO_SECURE
        )

    # ── 3. 下载文件 & 生成 CSV ────────────────────────────────────────────────
    csv_path = os.path.join(OUTPUT_DIR, 'summary.csv')
    csv_rows = []
    ok_count = err_count = skip_count = 0

    for row in rows:
        r = dict(zip(cols, row))

        reg_id       = r['reg_id']
        comp_name    = r['competition_name'] or ''
        proj_name    = r['project_name'] or ''
        group_type   = r['group_type'] or ''
        reg_status   = r['reg_status'] or ''
        inst_name    = r['institution_name'] or ''
        appl_name    = r['applicant_name'] or ''
        appl_phone   = r['applicant_phone'] or ''
        file_name    = r['file_name'] or ''
        file_url     = r['file_url'] or ''
        uploaded_at  = r['uploaded_at']

        # 本地保存目录：files/{reg_id}_{机构名}_{报名人}/
        folder_name = safe_name(f"{reg_id}_{inst_name}_{appl_name}")
        local_dir   = os.path.join(files_dir, folder_name)

        download_status = 'SKIPPED'
        local_path = ''

        if not SKIP_DOWNLOAD and file_url:
            os.makedirs(local_dir, exist_ok=True)
            # file_url 存的是 MinIO object key，可能带前导 /，去掉
            object_key = file_url.lstrip('/')
            local_path = os.path.join(local_dir, safe_name(file_name) or safe_name(object_key.split('/')[-1]))
            try:
                minio_client.fget_object(MINIO_BUCKET, object_key, local_path)
                download_status = 'OK'
                ok_count += 1
            except S3Error as e:
                download_status = f'ERROR: {e}'
                err_count += 1
                print(f'  !! 下载失败 reg={reg_id} file={file_url}: {e}')
            except Exception as e:
                download_status = f'ERROR: {e}'
                err_count += 1
                print(f'  !! 下载异常 reg={reg_id} file={file_url}: {e}')
        elif SKIP_DOWNLOAD:
            skip_count += 1
        else:
            skip_count += 1

        csv_rows.append({
            '报名ID':     reg_id,
            '竞赛名称':   comp_name,
            '项目名称':   proj_name,
            '组别':       group_type,
            '报名状态':   reg_status,
            '机构名称':   inst_name,
            '报名人姓名': appl_name,
            '报名人手机': appl_phone,
            '文件名':     file_name,
            'MinIO路径':  file_url,
            '上传时间':   uploaded_at,
            '下载状态':   download_status,
            '本地路径':   local_path,
        })

    # ── 4. 写 CSV ─────────────────────────────────────────────────────────────
    fieldnames = list(csv_rows[0].keys()) if csv_rows else []
    with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(csv_rows)

    # ── 5. 同时尝试写 Excel（可选，需要 openpyxl）────────────────────────────
    xlsx_path = os.path.join(OUTPUT_DIR, 'summary.xlsx')
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '缴费回执汇总'

        header_fill = PatternFill('solid', fgColor='4F81BD')
        header_font = Font(color='FFFFFF', bold=True)

        for col_idx, name in enumerate(fieldnames, 1):
            cell = ws.cell(row=1, column=col_idx, value=name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        for row_idx, row_data in enumerate(csv_rows, 2):
            for col_idx, key in enumerate(fieldnames, 1):
                val = row_data[key]
                if hasattr(val, 'strftime'):
                    val = val.strftime('%Y-%m-%d %H:%M:%S')
                ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else '')

        # 自动列宽
        for col in ws.columns:
            max_len = max((len(str(c.value or '')) for c in col), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

        wb.save(xlsx_path)
        print(f'Excel 已保存: {xlsx_path}')
    except ImportError:
        print('提示: 安装 openpyxl 可同时生成 Excel 文件（pip install openpyxl）')

    # ── 6. 汇总输出 ───────────────────────────────────────────────────────────
    print()
    print('=' * 60)
    print(f'导出完成！')
    print(f'  记录总数  : {len(csv_rows)}')
    if not SKIP_DOWNLOAD:
        print(f'  下载成功  : {ok_count}')
        print(f'  下载失败  : {err_count}')
        print(f'  无URL跳过 : {skip_count}')
    print(f'  CSV汇总   : {csv_path}')
    if not SKIP_DOWNLOAD:
        print(f'  文件目录  : {files_dir}')
    print('=' * 60)


if __name__ == '__main__':
    main()

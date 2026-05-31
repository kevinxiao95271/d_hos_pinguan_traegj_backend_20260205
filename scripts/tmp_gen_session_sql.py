import openpyxl, io, sys, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

path = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205\4.排程-20260526(1).xlsx'
wb = openpyxl.load_workbook(path, read_only=True)

# 根据场次名称判断打分表类型
def infer_form(session_code):
    if '课题达成' in session_code or 'QFD' in session_code:
        return 'QFD'
    elif '十大安全' in session_code:
        return 'NON_QCC'
    elif 'PDCA' in session_code:
        return 'QCC'
    elif '综合工具' in session_code:
        return 'QCC'
    elif '问题解决' in session_code:
        return 'QCC'
    elif '进阶' in session_code:
        return 'QCC'
    elif '基层' in session_code:
        return 'QCC'
    else:
        return 'QCC'

lines = []
lines.append("-- 现场竞赛分组数据 UPDATE（生成自 4.排程-20260526(1).xlsx）")
lines.append("-- 更新 registrations 表: final_session_date, final_session_code, final_session_order, final_score_form")
lines.append("")

total = 0
for shname in wb.sheetnames:
    # 跳过汇总 sheet
    if shname == '进阶组':
        continue

    # 解析日期前缀：6.3 / 6.4 / 6.5
    m = re.match(r'^(6\.[345])', shname)
    if not m:
        continue
    date_raw = m.group(1)
    date_map = {'6.3': '20260603', '6.4': '20260604', '6.5': '20260605'}
    session_date = date_map[date_raw]

    # 场次代码：去掉日期前缀
    session_code = shname[len(date_raw):].strip()

    score_form = infer_form(session_code)

    ws = wb[shname]
    rows = list(ws.iter_rows(values_only=True))

    lines.append(f"-- ===== {shname} =====")

    for row in rows[1:]:  # 跳过表头
        if not row or len(row) < 3:
            continue
        order_val = row[1]
        reg_id = row[2]

        # 跳过非数字顺序（如"交流时间"、"中餐时间"）
        if not isinstance(order_val, (int, float)):
            continue
        if not isinstance(reg_id, (int, float)):
            continue

        order = int(order_val)
        reg_id = int(reg_id)

        safe_code = session_code.replace("'", "''")
        lines.append(
            f"UPDATE registrations SET "
            f"final_session_date='{session_date}', "
            f"final_session_code='{safe_code}', "
            f"final_session_order={order}, "
            f"final_score_form='{score_form}' "
            f"WHERE id={reg_id};"
        )
        total += 1

lines.append("")
lines.append(f"-- 共 {total} 条")

sql = '\n'.join(lines)

out_path = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205\scripts\update_final_sessions_0527.sql'
with open(out_path, 'w', encoding='utf-8') as f:
    f.write(sql)

print(f"生成完成: {total} 条 UPDATE")
print(f"输出到: {out_path}")
print("\n=== 场次汇总 ===")
for shname in wb.sheetnames:
    if shname == '进阶组': continue
    m = re.match(r'^(6\.[345])', shname)
    if not m: continue
    date_raw = m.group(1)
    session_code = shname[len(date_raw):].strip()
    score_form = infer_form(session_code)
    ws = wb[shname]
    rows = list(ws.iter_rows(values_only=True))
    cnt = sum(1 for r in rows[1:] if r and len(r)>=3 and isinstance(r[1],(int,float)) and isinstance(r[2],(int,float)))
    print(f"  {shname[:35]:35s} → {score_form:8s} {cnt}个项目")

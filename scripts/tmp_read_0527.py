import os, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

root = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205'
files = os.listdir(root)

# 找两个目标文件
inst_file = None
proj_file = None
for f in files:
    try:
        name = f.encode('cp936', errors='replace').decode('cp936', errors='replace')
    except:
        name = f
    if '0527' in f and ('机构' in f or '\xd2\xbd' in f or 'inst' in f.lower() or f.startswith('\xd2')):
        inst_file = f
    if '0527' in f and ('项目' in f or '\xc4\xbf' in f or 'proj' in f.lower()):
        proj_file = f

print("根目录文件列表:")
for f in sorted(files):
    if '0527' in f:
        print(f"  [{repr(f)}]")

print("\n===== 尝试读取所有含0527的xlsx =====")
for f in sorted(files):
    if '0527' in f and f.endswith('.xlsx'):
        path = os.path.join(root, f)
        print(f"\n文件: {f}")
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
            for shname in wb.sheetnames:
                ws = wb[shname]
                print(f"  Sheet: {shname}")
                rows = list(ws.iter_rows(values_only=True))
                for i, row in enumerate(rows[:30]):
                    if any(c is not None for c in row):
                        print(f"    第{i+1}行: {row}")
        except Exception as e:
            print(f"  读取失败: {e}")

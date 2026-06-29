import openpyxl, re, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from collections import defaultdict
import os

# 找排程文件
base = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205'
schedule_file = None
for f in os.listdir(base):
    if '排程' in f and '0526' in f and f.endswith('.xlsx'):
        schedule_file = os.path.join(base, f)
        break
print(f"排程文件: {schedule_file}")

wb_s = openpyxl.load_workbook(schedule_file, data_only=True)
print(f"Sheets: {wb_s.sheetnames}")

# 先打印每个sheet前3行，看结构
for shname in wb_s.sheetnames[:5]:
    ws = wb_s[shname]
    rows = list(ws.iter_rows(values_only=True))
    print(f"\n--- Sheet: {shname} ---")
    for r in rows[:3]:
        print(r)

import openpyxl, re, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

path = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205\现场竞赛项目分组及评分表类型 - 发工程师(1).xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

# 收集所有项目的正确评分表
# 列结构不完全一致，需要找"项目编号"列和"评分表"列的index
id_to_form = {}

for shname in wb.sheetnames:
    m = re.match(r'^(6\.[345])', shname)
    if not m:
        continue
    ws = wb[shname]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue

    # 找表头行
    header = rows[0]
    try:
        id_col = next(i for i, h in enumerate(header) if h == '项目编号')
        form_col = next(i for i, h in enumerate(header) if h == '评分表')
    except StopIteration:
        print(f"  {shname}: 找不到列头，跳过")
        continue

    for r in rows[1:]:
        if not r or len(r) <= max(id_col, form_col):
            continue
        reg_id = r[id_col]
        form_val = r[form_col]
        if not isinstance(reg_id, (int, float)):
            continue
        reg_id = int(reg_id)
        if form_val == 'QCC':
            db_form = 'QCC'
        elif form_val == 'QFD':
            db_form = 'QFD'
        elif form_val == '非QCC':
            db_form = 'NON_QCC'
        else:
            continue  # 跳过空值或异常值
        id_to_form[reg_id] = db_form

print(f"Excel中共读到 {len(id_to_form)} 条项目")

# 按期望值分布
from collections import Counter
dist = Counter(id_to_form.values())
print(f"期望分布: {dict(dist)}")

# 生成查询 SQL：找出 final_score_form 与期望不符的记录
# 用 CASE WHEN 构建期望值表，再 WHERE 不等于
lines = []
lines.append("-- 查询 final_score_form 与正确值不符的项目")
lines.append("SELECT")
lines.append("    r.id,")
lines.append("    r.final_session_code,")
lines.append("    r.final_session_order,")
lines.append("    r.final_score_form AS 当前值,")
lines.append("    CASE r.id")
for reg_id, form in sorted(id_to_form.items()):
    lines.append(f"        WHEN {reg_id} THEN '{form}'")
lines.append("        ELSE NULL")
lines.append("    END AS 正确值")
lines.append("FROM registrations r")
lines.append("WHERE r.final_session_code IS NOT NULL")
lines.append("  AND r.final_score_form != CASE r.id")
for reg_id, form in sorted(id_to_form.items()):
    lines.append(f"        WHEN {reg_id} THEN '{form}'")
lines.append("        ELSE r.final_score_form")
lines.append("    END")
lines.append("ORDER BY r.final_session_code, r.final_session_order;")

sql = '\n'.join(lines)
out = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205\scripts\check_score_form.sql'
with open(out, 'w', encoding='utf-8') as f:
    f.write(sql)
print(f"\nSQL已保存到: {out}")
print("\n--- 汇总查询（看总数）---")
print("""SELECT
    CASE r.id""")
for reg_id, form in sorted(id_to_form.items()):
    print(f"        WHEN {reg_id} THEN '{form}'")
print("""        ELSE r.final_score_form
    END AS 正确值,
    r.final_score_form AS 当前值,
    COUNT(*) AS 数量
FROM registrations r
WHERE r.final_session_code IS NOT NULL
GROUP BY 正确值, 当前值
ORDER BY 正确值, 当前值;""")

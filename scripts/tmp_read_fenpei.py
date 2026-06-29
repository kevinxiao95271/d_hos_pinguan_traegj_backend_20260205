import openpyxl, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

path = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205\现场竞赛项目分组及评分表类型 - 发工程师(1).xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
print("Sheet列表:", wb.sheetnames)

for shname in wb.sheetnames:
    ws = wb[shname]
    rows = list(ws.iter_rows(values_only=True))
    non_empty = [r for r in rows if any(c is not None for c in r)]
    print(f"\n===== Sheet: {shname} ({len(non_empty)}行) =====")
    for i, r in enumerate(non_empty[:8]):
        print(f"  {i}: {r}")
    if len(non_empty) > 8:
        print(f"  ... 共{len(non_empty)}行")

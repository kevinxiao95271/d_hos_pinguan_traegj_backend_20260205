import openpyxl, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

path = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205\4.排程-20260526(1).xlsx'
wb = openpyxl.load_workbook(path, read_only=True)
print(f"Sheet列表: {wb.sheetnames}\n")

for shname in wb.sheetnames:
    ws = wb[shname]
    rows = list(ws.iter_rows(values_only=True))
    non_empty = [r for r in rows if any(c is not None for c in r)]
    print(f"===== Sheet: {shname} ({len(non_empty)}行有数据) =====")
    for i, row in enumerate(non_empty[:50]):
        print(f"  {i+1}: {row}")
    if len(non_empty) > 50:
        print(f"  ... 共 {len(non_empty)} 行，只显示前50行")
    print()

"""扫描汇报书DOCX的关键段落，找到方法论线索"""
import os, sys, re
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

BASE = r'D:\AiCode\cursor\d_hos_pinguan_traegj_backend_20260205\scripts\7error'

def get_folder(reg_id):
    return os.path.join(BASE, [d for d in os.listdir(BASE) if d.startswith(reg_id)][0])

def scan_docx_for_clues(reg_id, fname_part):
    """在文件名含fname_part的DOCX里找方法论线索"""
    folder = get_folder(reg_id)
    target = None
    for fn in os.listdir(folder):
        if fname_part in fn and fn.endswith('.docx') and not fn.startswith('~$'):
            target = fn
            break
    if not target:
        print(f'  [{reg_id}] 未找到含"{fname_part}"的DOCX')
        return
    fp = os.path.join(folder, target)
    print(f'\n[{reg_id}] {target}')
    try:
        doc = Document(fp)
        # 收集所有段落文字
        paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        # 收集所有表格的第一列
        table_first_col = []
        for t in doc.tables:
            for row in t.rows:
                if row.cells:
                    v = row.cells[0].text.strip()
                    if v:
                        table_first_col.append(v)

        all_text = '\n'.join(paras + table_first_col)

        # 找品管圈/PDCA/精益等方法论关键词
        keywords = ['品管圈', 'QCC', 'PDCA', 'FOCUS', '精益', 'A3', '根本原因', 'FMEA',
                    '六西格玛', '循证', '标杆', '平衡计分', '主题类型', '运用手法',
                    '跨部门', '数字化', '医疗质量安全', '改善就医']
        found = {}
        for kw in keywords:
            if kw in all_text:
                # 找前后上下文
                idx = all_text.find(kw)
                ctx = all_text[max(0,idx-40):idx+80].replace('\n',' ')
                found[kw] = ctx

        for kw, ctx in found.items():
            print(f'  [{kw}] ...{ctx}...')

        # 前几段落
        print('  前20段落:')
        for p in paras[:20]:
            print(f'    {p}')
    except Exception as e:
        print(f'  失败: {e}')

print('='*60)
scan_docx_for_clues('20260563', '汇报书')

print()
print('='*60)
scan_docx_for_clues('20260621', '精益A3')

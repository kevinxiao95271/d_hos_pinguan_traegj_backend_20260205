import openpyxl, re, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from collections import defaultdict

path = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205\现场竞赛项目分组及评分表类型 - 发工程师(1).xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

# session_code -> {QCC:n, NON_QCC:n, QFD:n, total:n}
session_stats = defaultdict(lambda: defaultdict(int))

for shname in wb.sheetnames:
    if not re.match(r'^6\.[345]', shname):
        continue
    ws = wb[shname]
    rows = list(ws.iter_rows(values_only=True))
    if not rows: continue
    header = rows[0]
    try:
        id_col   = next(i for i,h in enumerate(header) if h == '项目编号')
        form_col = next(i for i,h in enumerate(header) if h == '评分表')
    except StopIteration:
        print(f"跳过 {shname}: 找不到列")
        continue

    # 从 sheet 名提取场次名（去掉前缀 "6.X"）
    sess_label = re.sub(r'^6\.[345]', '', shname).strip()

    for r in rows[1:]:
        if not r or len(r) <= max(id_col, form_col): continue
        rid = r[id_col]
        fv  = r[form_col]
        if not isinstance(rid, (int, float)): continue
        if fv == 'QCC':     form = 'QCC'
        elif fv == 'QFD':   form = 'QFD'
        elif fv == '非QCC': form = 'NON_QCC'
        else: continue
        session_stats[sess_label][form] += 1
        session_stats[sess_label]['total'] += 1

total_projects = sum(d['total'] for d in session_stats.values())
print(f"共 {total_projects} 个项目，{len(session_stats)} 个场次\n")

for sess in sorted(session_stats.keys()):
    d = session_stats[sess]
    parts = []
    for k in ['QCC','NON_QCC','QFD']:
        if d[k]: parts.append(f"{k}x{d[k]}")
    print(f"{sess}  {d['total']}项  {' / '.join(parts)}")

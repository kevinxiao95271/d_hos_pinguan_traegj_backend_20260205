import os
import re
import io
import sys
from typing import Dict, List, Any

import pymysql

try:
    import openpyxl
except ImportError:
    print("missing openpyxl, please install: pip install openpyxl")
    raise

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = r"D:\AiCode\cursor\d_hos_pinguan_traegj_backend_20260205"
DB = dict(
    host="gz-cdb-bq7gk3k5.sql.tencentcdb.com",
    port=63606,
    user="root",
    password="Yiguo9527_",
    db="d_hos_pinguan_traegj_20260205",
    charset="utf8mb4",
)


def norm(s: Any) -> str:
    if s is None:
        return ""
    return str(s).strip().replace("\n", "").replace(" ", "")


def find_file() -> str:
    for n in os.listdir(ROOT):
        if n.lower().endswith(".xlsx") and "(3)" in n:
            return os.path.join(ROOT, n)
    raise FileNotFoundError("xlsx with (3) not found")


def split_multi(v: str) -> List[str]:
    if not v:
        return []
    arr = [x.strip() for x in re.split(r"[、,，;；/]+", v) if x.strip()]
    return arr


def main():
    path = find_file()
    print(f"file: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    print(f"sheet: {ws.title}, rows={ws.max_row}, cols={ws.max_column}")

    # 识别表头行：包含“姓名/所属机构/手机号”等关键字
    header_row = None
    headers = []
    for r in range(1, min(ws.max_row, 20) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        joined = "|".join(vals)
        if ("姓名" in joined and ("所属机构" in joined or "机构" in joined) and ("手机" in joined or "电话" in joined)):
            header_row = r
            headers = vals
            break
    if not header_row:
        header_row = 1
        headers = [norm(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]

    print(f"header_row={header_row}")
    for i, h in enumerate(headers, start=1):
        if h:
            print(f"  col{i}: {h}")

    def find_col(*keys):
        for idx, h in enumerate(headers, start=1):
            hh = h
            if not hh:
                continue
            for k in keys:
                if k in hh:
                    return idx
        return None

    c_name = find_col("姓名")
    c_inst = find_col("所属机构", "机构")
    c_phone = find_col("手机号", "联系电话", "联系方式", "手机")
    c_gender = find_col("性别")
    c_title = find_col("职称")
    c_position = find_col("职务")
    c_idno = find_col("身份证号")
    c_idimg = find_col("身份证正反面", "身份证照片")
    c_bank = find_col("开户银行")
    c_card = find_col("银行卡号")
    c_bg = find_col("专业背景")
    c_tools = find_col("熟悉的品管工具", "品管工具")
    c_topics = find_col("擅长评审主题", "评审主题方向", "擅长主题")

    print("\nresolved columns:")
    print(f"name={c_name}, inst={c_inst}, phone={c_phone}, gender={c_gender}, title={c_title}, position={c_position}")
    print(f"idno={c_idno}, idimg={c_idimg}, bank={c_bank}, card={c_card}, bg={c_bg}, tools={c_tools}, topics={c_topics}")

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = norm(ws.cell(r, c_name).value) if c_name else ""
        inst = norm(ws.cell(r, c_inst).value) if c_inst else ""
        phone = norm(ws.cell(r, c_phone).value) if c_phone else ""
        if not name and not phone:
            continue
        row = {
            "row": r,
            "name": name,
            "inst": inst,
            "phone": phone,
            "gender": norm(ws.cell(r, c_gender).value) if c_gender else "",
            "title": norm(ws.cell(r, c_title).value) if c_title else "",
            "position": norm(ws.cell(r, c_position).value) if c_position else "",
            "idno": norm(ws.cell(r, c_idno).value) if c_idno else "",
            "idimg": norm(ws.cell(r, c_idimg).value) if c_idimg else "",
            "bank": norm(ws.cell(r, c_bank).value) if c_bank else "",
            "card": norm(ws.cell(r, c_card).value) if c_card else "",
            "bg": norm(ws.cell(r, c_bg).value) if c_bg else "",
            "tools": norm(ws.cell(r, c_tools).value) if c_tools else "",
            "topics": norm(ws.cell(r, c_topics).value) if c_topics else "",
        }
        rows.append(row)

    print(f"\nvalid candidate rows={len(rows)}")

    conn = pymysql.connect(**DB)
    cur = conn.cursor()

    ok = 0
    name_not_found = 0
    inst_mismatch = 0
    multi_users = 0
    issue_rows = []
    insert_fit_issue = []

    for x in rows:
        # 先按姓名+角色查
        cur.execute(
            """
            SELECT ua.id, ua.name, ua.phone, i.name
            FROM user_accounts ua
            LEFT JOIN institutions i ON ua.institution_id = i.id
            WHERE ua.role='REVIEWER' AND ua.name=%s
            """,
            (x["name"],),
        )
        found = cur.fetchall()
        if not found:
            name_not_found += 1
            issue_rows.append((x["row"], x["name"], x["phone"], x["inst"], "name_not_found"))
            continue

        # 如果手机号可用，优先按手机号定位
        picked = None
        if x["phone"]:
            for f in found:
                if norm(f[2]) == x["phone"]:
                    picked = f
                    break
        if not picked:
            if len(found) > 1:
                multi_users += 1
                issue_rows.append((x["row"], x["name"], x["phone"], x["inst"], "multi_reviewer_same_name"))
            picked = found[0]

        db_inst = norm(picked[3])
        excel_inst = norm(x["inst"])
        if excel_inst and db_inst and excel_inst != db_inst:
            inst_mismatch += 1
            issue_rows.append((x["row"], x["name"], x["phone"], f"{excel_inst} != {db_inst}", "institution_mismatch"))
        else:
            ok += 1

        # 能否落 reviewer_profiles（长度检查）
        errs = []
        if len(x["gender"]) > 8:
            errs.append("gender>8")
        if len(x["position"]) > 64:
            errs.append("position>64")
        if len(x["idno"]) > 64:
            errs.append("id_number>64")
        if len(x["bank"]) > 128:
            errs.append("bank_name>128")
        if len(x["card"]) > 128:
            errs.append("bank_card_no>128")
        # json text 基本不会超；这里只检查是否可拆分
        _ = split_multi(x["bg"])
        _ = split_multi(x["tools"])
        _ = split_multi(x["topics"])
        if errs:
            insert_fit_issue.append((x["row"], x["name"], ",".join(errs)))

    print("\n=== compare summary ===")
    print(f"matched_or_basically_ok={ok}")
    print(f"name_not_found={name_not_found}")
    print(f"institution_mismatch={inst_mismatch}")
    print(f"multi_reviewer_same_name={multi_users}")

    print("\n=== issues (top 30) ===")
    for i, r in enumerate(issue_rows[:30], start=1):
        print(f"{i:02d}. row={r[0]} name={r[1]} phone={r[2]} detail={r[3]} type={r[4]}")

    print("\n=== reviewer_profiles insert-fit issues ===")
    if not insert_fit_issue:
        print("none")
    else:
        for r in insert_fit_issue[:30]:
            print(f"row={r[0]} name={r[1]} err={r[2]}")

    print("\n=== map readiness ===")
    missing_cols = []
    if not c_name:
        missing_cols.append("姓名")
    if not c_inst:
        missing_cols.append("所属机构")
    if not c_phone:
        missing_cols.append("手机号")
    if missing_cols:
        print("critical_missing_columns:", ",".join(missing_cols))
    else:
        print("critical columns present")
    print("optional columns present:",
          {
              "gender": bool(c_gender), "title": bool(c_title), "position": bool(c_position),
              "idno": bool(c_idno), "idimg": bool(c_idimg), "bank": bool(c_bank), "card": bool(c_card),
              "background": bool(c_bg), "tools": bool(c_tools), "topics": bool(c_topics),
          })

    conn.close()


if __name__ == "__main__":
    main()


import openpyxl, re, io, sys, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from collections import defaultdict

base = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205'

# 弃赛
wb_w = openpyxl.load_workbook(os.path.join(base, '弃赛项目汇总(1).xlsx'), data_only=True)
withdrawn = set()
for r in wb_w['Sheet1'].iter_rows(values_only=True):
    rid = r[4]
    if isinstance(rid, (int, float)) and rid > 20000000:
        withdrawn.add(int(rid))

# 评分表类型 Excel
wb_f = openpyxl.load_workbook(os.path.join(base, '现场竞赛项目分组及评分表类型 - 发工程师(1).xlsx'), data_only=True)
session_stats = defaultdict(lambda: defaultdict(int))

for shname in wb_f.sheetnames:
    if not re.match(r'^6\.[345]', shname): continue
    ws = wb_f[shname]
    rows = list(ws.iter_rows(values_only=True))
    if not rows: continue
    header = rows[0]
    try:
        id_col   = next(i for i,h in enumerate(header) if h == '项目编号')
        form_col = next(i for i,h in enumerate(header) if h == '评分表')
    except StopIteration:
        continue
    sess_label = re.sub(r'^6\.[345]', '', shname).strip()
    for r in rows[1:]:
        if not r or len(r) <= max(id_col, form_col): continue
        rid = r[id_col]
        fv  = r[form_col]
        if not isinstance(rid, (int, float)): continue
        rid = int(rid)
        if rid in withdrawn: continue   # 排除弃赛
        if fv == 'QCC':     form = 'QCC'
        elif fv == 'QFD':   form = 'QFD'
        elif fv == '非QCC': form = 'NON_QCC'
        else: continue
        session_stats[sess_label][form] += 1
        session_stats[sess_label]['total'] += 1

total = sum(d['total'] for d in session_stats.values())
print(f"排除弃赛后合计: {total} 项，{len(session_stats)} 个场次\n")

# 用户提供的预期值
expected = {
    '基层组1组（一楼活动中心）':         {'total':29,'QCC':10,'NON_QCC':18,'QFD':1},
    '基层组2组（锦绣厅）':               {'total':28,'QCC':18,'NON_QCC':10},
    '综合组-PDCA专场1（二楼名仕厅）':    {'total':24,'NON_QCC':24},
    '综合组-综合工具专场1（三楼锦兰厅）':{'total':17,'NON_QCC':17},
    '综合组-课题达成及QFD专场2改1（三楼开元B厅）':{'total':28,'QCC':23,'QFD':5},
    '综合组-问题解决型专场1（三楼开元A厅）':{'total':22,'QCC':22},
    '进阶组1组（三楼萧然厅）':           {'total':20,'QCC':12,'NON_QCC':5,'QFD':3},
    '基层组3组（一楼活动中心）':         {'total':28,'QCC':15,'NON_QCC':13},
    '综合组-PDCA专场2（二楼名仕厅）':    {'total':24,'NON_QCC':24},
    '综合组-十大安全目标专场1（三楼锦绣厅）':{'total':24,'QCC':17,'NON_QCC':7},
    '综合组-综合工具专场2（三楼锦兰厅）':{'total':21,'NON_QCC':21},
    '综合组-课题达成及QFD专场2（三楼开元B厅）':{'total':27,'QCC':21,'QFD':6},
    '综合组-问题解决型专场2（三楼开元A厅）':{'total':25,'QCC':25},
    '进阶组 2组（三楼萧然厅）':          {'total':20,'QCC':10,'NON_QCC':6,'QFD':4},
    '基层组4组（一楼活动中心）':         {'total':28,'QCC':13,'NON_QCC':14,'QFD':1},
    '综合组-PDCA专场3（二楼名仕厅）':    {'total':24,'NON_QCC':24},
    '综合组-十大安全目标专场2（三楼锦绣厅）':{'total':22,'QCC':10,'NON_QCC':12},
    '综合组-课题达成及QFD专场3（三楼开元B厅）':{'total':27,'QCC':24,'QFD':3},
    '综合组-问题解决型专场3（三楼锦兰厅）':{'total':24,'QCC':24},
    '综合组-问题解决型专场4(三楼开元A厅)':{'total':25,'QCC':25},
    '进阶组3组（三楼萧然厅）':           {'total':20,'QCC':11,'NON_QCC':6,'QFD':3},
}

# 注意：Excel里课题达成QFD专场2叫"课题达成及QFD专场1改2"，需要对应
rename_map = {
    '综合组-课题达成及QFD专场1改2（三楼开元B厅）': '综合组-课题达成及QFD专场2（三楼开元B厅）',
}

all_ok = True
for sess in sorted(session_stats.keys()):
    d = session_stats[sess]
    # 查预期（可能需要rename）
    exp_key = rename_map.get(sess, sess)
    exp = expected.get(exp_key)
    actual_parts = ' / '.join(f"{k}x{d[k]}" for k in ['QCC','NON_QCC','QFD'] if d[k])
    if exp is None:
        print(f"[?无预期] {sess}  {d['total']}项  {actual_parts}")
        continue
    ok = (d['total'] == exp.get('total',0)
          and d['QCC'] == exp.get('QCC',0)
          and d['NON_QCC'] == exp.get('NON_QCC',0)
          and d['QFD'] == exp.get('QFD',0))
    tag = 'OK' if ok else 'DIFF'
    if not ok: all_ok = False
    exp_parts = ' / '.join(f"{k}x{exp[k]}" for k in ['QCC','NON_QCC','QFD'] if exp.get(k,0))
    if ok:
        print(f"[{tag}] {sess}  {d['total']}项  {actual_parts}")
    else:
        print(f"[{tag}] {sess}")
        print(f"       实际: {d['total']}项  {actual_parts}")
        print(f"       预期: {exp.get('total',0)}项  {exp_parts}")

print(f"\n{'全部一致' if all_ok else '存在差异，见上方 [DIFF]'}")

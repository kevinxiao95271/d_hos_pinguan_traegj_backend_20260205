import os, sys, re
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

fp = r'D:\AiCode\cursor\d_hos_pinguan_traegj_backend_20260205\scripts\7error\20260893_杭州市富阳区第三人民医院\材料2.基于PDCA循环提高严重精神障碍患者随访信息录入规范率成果报告书.docx'

doc = Document(fp)
print(f'表格数: {len(doc.tables)}, 段落数: {len(doc.paragraphs)}')

# 找所有表格，看是否有主题类型等字段
for ti, table in enumerate(doc.tables):
    rows = []
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        seen = []
        for c in cells:
            if not seen or c != seen[-1]:
                seen.append(c)
        line = ' | '.join(seen)
        if line.strip():
            rows.append(line)
    full = '\n'.join(rows)
    # 检查是否包含关键字段
    keywords = ['主题类型','运用手法','跨部门','数字化','改善就医','医疗质量安全','关键词','平均工作年资','平均年龄']
    found = [k for k in keywords if k in full]
    if found:
        print(f'\n★ 表格{ti+1} 包含: {found}')
        print(full[:2000])

# 段落里找关键词
print('\n=== 相关段落 ===')
for p in doc.paragraphs:
    t = p.text.strip()
    if any(k in t for k in ['主题类型','运用手法','跨部门','数字化','医疗质量','改善就医','关键词','品管圈','PDCA','精益']):
        print(f'  {t}')

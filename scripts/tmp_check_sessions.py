import openpyxl, bcrypt, io, sys, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# ---------- 1. 读出所有真实 session_code ----------
schedule_path = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205\4.排程-20260526(1).xlsx'
wb = openpyxl.load_workbook(schedule_path, read_only=True)
real_codes = {}          # 简化名 -> 真实 session_code
all_real_codes = []
for shname in wb.sheetnames:
    m = re.match(r'^(6\.[345])', shname)
    if not m:
        continue
    date_raw = m.group(1)
    session_code = shname[len(date_raw):].strip()
    if session_code not in [c for c, _ in all_real_codes]:
        all_real_codes.append((session_code, shname))

print("=== 所有真实 session_code ===")
for code, sh in all_real_codes:
    print(f"  [{sh}] -> {code}")
print()

# ---------- 2. 读工作人员权限文件 ----------
staff_path = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205\工作人员权限开通(1).xlsx'
wb2 = openpyxl.load_workbook(staff_path, read_only=True)
ws2 = wb2.active
rows2 = list(ws2.iter_rows(values_only=True))
print("=== 工作人员权限文件内容 ===")
for r in rows2:
    print(r)

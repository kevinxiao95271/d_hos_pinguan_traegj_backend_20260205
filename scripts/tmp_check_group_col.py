import openpyxl, re, io, sys, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

base = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205'
wb = openpyxl.load_workbook(os.path.join(base, '现场竞赛项目分组及评分表类型 - 发工程师(1).xlsx'), data_only=True)

for shname in wb.sheetnames:
    if not re.match(r'^6\.[345]', shname): continue
    ws = wb[shname]
    rows = list(ws.iter_rows(values_only=True))
    if not rows: continue
    header = rows[0]
    # 找组别列
    group_col = next((i for i,h in enumerate(header) if h == '组别'), None)
    if group_col is None:
        print(f"{shname}: 无组别列，header={header}")
        continue
    # 取前3个数据行的组别值
    vals = []
    for r in rows[1:4]:
        if r and len(r) > group_col:
            vals.append(r[group_col])
    print(f"{shname}: {vals}")

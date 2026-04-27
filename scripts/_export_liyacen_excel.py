"""
导出李雅岑所有打分操作记录到 Excel（每次草稿保存/提交都是一行）
"""
import struct, sys, io, datetime, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import defaultdict
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

TARGET_REVIEWER_ID = 703
MAGIC = b'\xfebin'

TABLE_MAP_EVENT = 19
WRITE_ROWS_V1, UPDATE_ROWS_V1, DELETE_ROWS_V1 = 23, 24, 25
WRITE_ROWS_V2, UPDATE_ROWS_V2, DELETE_ROWS_V2 = 30, 31, 32
ROW_WRITE  = {WRITE_ROWS_V1, WRITE_ROWS_V2}
ROW_UPDATE = {UPDATE_ROWS_V1, UPDATE_ROWS_V2}
ROW_DELETE = {DELETE_ROWS_V1, DELETE_ROWS_V2}
ROW_EVENTS = ROW_WRITE | ROW_UPDATE | ROW_DELETE

TYPE_LONGLONG=8; TYPE_DOUBLE=5; TYPE_VARCHAR=15; TYPE_VAR_STR=253
TYPE_BLOB=252; TYPE_DATETIME2=18; TYPE_LONG=3; TYPE_SHORT=2
TYPE_TINY=1; TYPE_ENUM=247; TYPE_SET=248; TYPE_STRING=254

def read_lenenc(data, pos):
    if pos >= len(data): return 0, pos
    b = data[pos]
    if b < 0xfb:   return b, pos+1
    if b == 0xfc:  return struct.unpack_from('<H', data, pos+1)[0], pos+3
    if b == 0xfd:  return int.from_bytes(data[pos+1:pos+4],'little'), pos+4
    return struct.unpack_from('<Q', data, pos+1)[0], pos+9

def parse_table_map(body):
    pos = 0
    tid = int.from_bytes(body[pos:pos+6],'little'); pos = 8
    db_len = body[pos]; pos+=1
    db = body[pos:pos+db_len].decode('utf-8','replace'); pos+=db_len+1
    tl = body[pos]; pos+=1
    tbl = body[pos:pos+tl].decode('utf-8','replace'); pos+=tl+1
    col_count, pos = read_lenenc(body, pos)
    col_types = list(body[pos:pos+col_count]); pos+=col_count
    meta_len, pos = read_lenenc(body, pos)
    meta_raw = body[pos:pos+meta_len]
    metas = {}; mp = 0
    for ci, ct in enumerate(col_types):
        if ct in (TYPE_VARCHAR,TYPE_VAR_STR,TYPE_STRING):
            if mp+2<=len(meta_raw): metas[ci]=struct.unpack_from('<H',meta_raw,mp)[0]; mp+=2
        elif ct in (TYPE_DOUBLE,TYPE_BLOB,TYPE_DATETIME2,TYPE_ENUM,TYPE_SET):
            if mp<len(meta_raw): metas[ci]=meta_raw[mp]; mp+=1
    return tid, db, tbl, col_types, metas

def read_col(data, pos, ctype, meta):
    try:
        if ctype==TYPE_LONGLONG: return struct.unpack_from('<q',data,pos)[0], pos+8
        if ctype==TYPE_LONG:     return struct.unpack_from('<i',data,pos)[0], pos+4
        if ctype==TYPE_SHORT:    return struct.unpack_from('<h',data,pos)[0], pos+2
        if ctype==TYPE_TINY:     return data[pos], pos+1
        if ctype==TYPE_DOUBLE:   return struct.unpack_from('<d',data,pos)[0], pos+8
        if ctype in (TYPE_VARCHAR,TYPE_VAR_STR,TYPE_STRING):
            sb = 2 if meta>255 else 1
            sl = int.from_bytes(data[pos:pos+sb],'little'); pos+=sb
            return data[pos:pos+sl].decode('utf-8','replace'), pos+sl
        if ctype==TYPE_BLOB:
            lb = meta if meta else 2
            sl = int.from_bytes(data[pos:pos+lb],'little'); pos+=lb
            return data[pos:pos+sl].decode('utf-8','replace'), pos+sl
        if ctype==TYPE_DATETIME2:
            raw = int.from_bytes(data[pos:pos+5],'big')
            fb = (meta+1)//2 if meta else 0
            pos2 = pos+5+fb
            ym=(raw>>22)&0x7FFF; yr,mo=ym//13,ym%13
            dy=(raw>>17)&0x1F; hr=(raw>>12)&0x1F
            mn=(raw>>6)&0x3F; sc=raw&0x3F
            try: return datetime.datetime(yr,mo,dy,hr,mn,sc), pos2
            except: return None, pos2
        if ctype in (TYPE_ENUM,TYPE_SET):
            nb=meta if meta else 1
            return int.from_bytes(data[pos:pos+nb],'little'), pos+nb
        return None, pos
    except: return None, pos+1

def parse_rows(body, etype, col_types, metas):
    rows=[]; pos=8
    if etype in (WRITE_ROWS_V2,UPDATE_ROWS_V2,DELETE_ROWS_V2):
        extra=struct.unpack_from('<H',body,pos)[0]; pos+=extra
    col_count, pos = read_lenenc(body, pos)
    bm_len=(col_count+7)//8
    pos+=bm_len
    if etype in ROW_UPDATE: pos+=bm_len
    while pos<len(body):
        null_bm=body[pos:pos+bm_len]; pos+=bm_len
        if etype in ROW_UPDATE:
            vals_b={}
            for ci,ct in enumerate(col_types):
                bb=ci//8; bi=ci%8
                if bb<len(null_bm) and (null_bm[bb]>>bi)&1: vals_b[ci]=None; continue
                v,pos=read_col(body,pos,ct,metas.get(ci,0)); vals_b[ci]=v
            null_bm=body[pos:pos+bm_len]; pos+=bm_len
        vals={}
        for ci,ct in enumerate(col_types):
            bb=ci//8; bi=ci%8
            if bb<len(null_bm) and (null_bm[bb]>>bi)&1: vals[ci]=None; continue
            v,pos=read_col(body,pos,ct,metas.get(ci,0))
            if isinstance(v,str) and v.startswith('ERR'): break
            vals[ci]=v
        if vals: rows.append(vals)
        else: break
    return rows

def scan_binlog(files, target_tables):
    table_map={}
    result={t:[] for t in target_tables}
    for fpath in files:
        with open(fpath,'rb') as f:
            if f.read(4)!=MAGIC: continue
            while True:
                hdr=f.read(19)
                if len(hdr)<19: break
                ts,etype,sid,elen,lpos,flags=struct.unpack('<IBIIIH',hdr)
                body=f.read(elen-19)
                if len(body)<elen-19: break
                dt=datetime.datetime.fromtimestamp(ts) if ts else None
                if etype==TABLE_MAP_EVENT:
                    try:
                        tid,db,tbl,ctypes,cmetas=parse_table_map(body)
                        table_map[tid]=(tbl,ctypes,cmetas)
                    except: pass
                elif etype in ROW_EVENTS:
                    tid=int.from_bytes(body[0:6],'little')
                    info=table_map.get(tid)
                    if not info or info[0] not in target_tables: continue
                    tbl,ctypes,cmetas=info
                    op='INSERT' if etype in ROW_WRITE else ('UPDATE' if etype in ROW_UPDATE else 'DELETE')
                    try:
                        for row in parse_rows(body,etype,ctypes,cmetas):
                            result[tbl].append((op,dt,row))
                    except: pass
    return result

BINLOG_FILES=['binlog.000003','binlog.000004']
print("扫描 binlog...")
data=scan_binlog(BINLOG_FILES,{'review_tasks','review_scores'})

# 步骤1: 找李雅岑 task_id
liyacen_tasks={}
for op,dt,vals in data['review_tasks']:
    if vals.get(3)==TARGET_REVIEWER_ID and vals.get(0):
        tid=vals.get(0)
        if tid not in liyacen_tasks:
            liyacen_tasks[tid]={'reg_id':vals.get(2)}

print(f"找到 task_id {len(liyacen_tasks)} 个: {sorted(liyacen_tasks.keys())}")

# 步骤2: 收集所有打分操作（每次都保留）
all_records=[]
for op,dt,vals in data['review_scores']:
    task_id=vals.get(1)
    if task_id not in liyacen_tasks: continue
    reg_id=liyacen_tasks[task_id]['reg_id']
    all_records.append({
        'task_id':    task_id,
        'reg_id':     reg_id,
        'op':         op,
        'time':       dt,
        'plan':       vals.get(2),
        'problem':    vals.get(3),
        'action':     vals.get(4),
        'success':    vals.get(5),
        'review':     vals.get(6),
        'operation':  vals.get(7),
        'presentation':vals.get(8),
        'total':      vals.get(9),
        'highlight':  vals.get(10),
        'weakness':   vals.get(11),
        'submitted_at':vals.get(12),
    })

all_records.sort(key=lambda x:(x['task_id'], x['time'] or datetime.datetime.min))
print(f"共 {len(all_records)} 条操作记录")

# 步骤3: 写 Excel
wb=openpyxl.Workbook()
ws=wb.active
ws.title='李雅岑打分记录'

# 样式
hdr_fill=PatternFill('solid',fgColor='1F4E79')
hdr_font=Font(bold=True,color='FFFFFF',size=10)
del_fill=PatternFill('solid',fgColor='FFE0E0')
sub_fill=PatternFill('solid',fgColor='E2EFDA')
thin=Side(style='thin',color='CCCCCC')
border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical='top')
center=Alignment(horizontal='center',vertical='top')

headers=['task_id','报名ID','操作类型','操作时间',
         '计划拟定','问题解析','对策拟定','成果','检讨改进','活动运作','报告呈现','总分',
         '亮点','不足','提交时间']
col_widths=[10,12,10,20,9,9,9,9,9,9,9,9,50,50,20]

for ci,(h,w) in enumerate(zip(headers,col_widths),1):
    cell=ws.cell(1,ci,h)
    cell.font=hdr_font; cell.fill=hdr_fill
    cell.alignment=center; cell.border=border
    ws.column_dimensions[cell.column_letter].width=w

ws.row_dimensions[1].height=20

for ri,r in enumerate(all_records,2):
    row_data=[
        r['task_id'], r['reg_id'], r['op'],
        r['time'].strftime('%Y-%m-%d %H:%M:%S') if r['time'] else '',
        r['plan'],r['problem'],r['action'],r['success'],
        r['review'],r['operation'],r['presentation'],r['total'],
        r['highlight'] or '',
        r['weakness'] or '',
        r['submitted_at'].strftime('%Y-%m-%d %H:%M:%S') if r['submitted_at'] else '',
    ]
    fill=None
    if r['op']=='DELETE': fill=del_fill
    elif r['submitted_at'] and r['op'] in ('INSERT','UPDATE'): fill=sub_fill

    for ci,v in enumerate(row_data,1):
        cell=ws.cell(ri,ci,v)
        cell.border=border
        cell.alignment=wrap
        if fill: cell.fill=fill
        if ci in range(5,13) and v is not None:
            try: cell.value=round(float(v),1)
            except: pass

    ws.row_dimensions[ri].height=60

# 冻结首行
ws.freeze_panes='A2'

# 第二个 sheet：每个 task 的最终状态（最后一次 INSERT/UPDATE）
ws2=wb.create_sheet('最终打分汇总')
for ci,(h,w) in enumerate(zip(headers,col_widths),1):
    cell=ws2.cell(1,ci,h)
    cell.font=hdr_font; cell.fill=hdr_fill
    cell.alignment=center; cell.border=border
    ws2.column_dimensions[cell.column_letter].width=w
ws2.row_dimensions[1].height=20

# 每个 task 取最后一次有效的 INSERT/UPDATE
from collections import OrderedDict
final=OrderedDict()
for r in all_records:
    if r['op'] in ('INSERT','UPDATE'):
        final[r['task_id']]=r

for ri,(tid,r) in enumerate(sorted(final.items()),2):
    row_data=[
        r['task_id'],r['reg_id'],r['op'],
        r['time'].strftime('%Y-%m-%d %H:%M:%S') if r['time'] else '',
        r['plan'],r['problem'],r['action'],r['success'],
        r['review'],r['operation'],r['presentation'],r['total'],
        r['highlight'] or '',r['weakness'] or '',
        r['submitted_at'].strftime('%Y-%m-%d %H:%M:%S') if r['submitted_at'] else '',
    ]
    fill=sub_fill if r['submitted_at'] else None
    for ci,v in enumerate(row_data,1):
        cell=ws2.cell(ri,ci,v)
        cell.border=border; cell.alignment=wrap
        if fill: cell.fill=fill
        if ci in range(5,13) and v is not None:
            try: cell.value=round(float(v),1)
            except: pass
    ws2.row_dimensions[ri].height=80

ws2.freeze_panes='A2'

out='scripts/李雅岑书审打分记录.xlsx'
wb.save(out)
print(f"\n✅ 已导出: {out}")
print(f"   Sheet1【李雅岑打分记录】: {len(all_records)} 行（每次操作一行，红色=DELETE，绿色=已提交）")
print(f"   Sheet2【最终打分汇总】:   {len(final)} 行（每个task最后一次有效打分）")

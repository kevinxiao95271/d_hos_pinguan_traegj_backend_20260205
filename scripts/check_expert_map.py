import openpyxl, os, sys
sys.stdout.reconfigure(encoding='utf-8')
BASE = r'd:\AiCode\cursor\d_hos_pinguan_traegj_backend_20260205'
wb = openpyxl.load_workbook(os.path.join(BASE, '书审专家单位(1)-0410.xlsx'), read_only=True)
ws = wb['书审']

print('=== 专家文件前8行原始内容 ===')
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i < 8:
        print(f'行{i+1}: {[str(v) if v is not None else "" for v in row]}')

print()
print('=== 所有专家分组编码及首个专家 ===')
cur = None
first_of_group = {}
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i < 2: continue
    if row[0]:
        cur = str(row[0]).strip()
    if cur and cur not in first_of_group and row[3]:
        first_of_group[cur] = str(row[3]).strip()

for g, name in first_of_group.items():
    print(f'  {g}: 首个专家={name}')

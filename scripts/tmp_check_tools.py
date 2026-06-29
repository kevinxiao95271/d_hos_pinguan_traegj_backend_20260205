import openpyxl, re, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

path = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205\4.排程-20260526(1).xlsx'
# data_only=True 读取公式缓存值
wb = openpyxl.load_workbook(path, data_only=True)

for shname in wb.sheetnames:
    m = re.match(r'^(6\.[345])', shname)
    if not m:
        continue
    ws = wb[shname]
    rows = list(ws.iter_rows(values_only=True))
    data_rows = [r for r in rows[1:] if r and len(r) >= 6
                 and isinstance(r[1], (int, float)) and isinstance(r[2], (int, float))]
    if not data_rows:
        continue

    # 统计最后列的唯一值
    tools = {}
    for r in data_rows:
        tool = r[5] if len(r) > 5 else None
        tools[tool] = tools.get(tool, 0) + 1

    date_raw = m.group(1)
    session_code = shname[len(date_raw):].strip()
    print(f"\n{session_code}")
    print(f"  项目数: {len(data_rows)}")
    print(f"  运用工具分布: {tools}")

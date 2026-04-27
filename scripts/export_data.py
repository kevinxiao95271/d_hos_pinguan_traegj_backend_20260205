"""
导出报名完整数据到 Excel
用法：python export_data.py [数量]
      python export_data.py 10    ← 前10条测试
      python export_data.py       ← 全量
注：applicant_phone / institution_city 接口未返回，留空
"""
import requests, sys, os
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

BASE  = 'http://zkjb.zjmss.org.cn'
LIMIT = int(sys.argv[1]) if len(sys.argv) > 1 else None

GROUP_MAP = {'ADVANCED': '进阶组', 'COMPREHENSIVE': '综合组', 'BASIC': '基层组'}

# 英文列名 → 中文列名
COL_ZH = {
    'registration_id':          '报名ID',
    'project_name':             '项目名称',
    'status':                   '报名状态',
    'submitted_at':             '提交时间',
    'group_type':               '组别代码',
    'group_type_label':         '组别',
    'group_code':               '组别编号',
    'created_at':               '创建时间',
    'applicant_name':           '报名人姓名',
    'applicant_phone':          '报名人手机',
    'institution_name':         '机构名称',
    'institution_level':        '机构级别',
    'institution_uscc':         '统一社会信用代码',
    'institution_region':       '所在区域',
    'institution_city':         '所在城市',
    'activity_theme':           '活动主题',
    'keywords':                 '关键词',
    'avg_work_years':           '平均工龄（年）',
    'avg_age':                  '平均年龄',
    'cross_department':         '是否跨部门',
    'related_to_digital_ai':    '是否数字化/人工智能应用相关',
    'subject_type_code':        '主题类型代码',
    'subject_type_label':       '主题类型',
    'subject_type_other':       '主题类型（其他说明）',
    'method_code':              '运用手法代码',
    'method_label':             '运用手法',
    'method_other':             '运用手法（其他说明）',
    'quality_topic_code':       '医疗质量安全主题代码',
    'quality_topic_label':      '医疗质量安全主题',
    'quality_topic_other':      '医疗质量安全主题（其他说明）',
    'experience_improve_code':  '改善就医感受类型代码',
    'experience_improve_label': '改善就医感受类型',
    'experience_improve_other': '改善就医感受类型（其他说明）',
    'summary_theme':            '摘要主题',
    'plan':                     '现状分析/计划',
    'problem':                  '问题把握',
    'action':                   '对策实施',
    'success':                  '成效',
    'discussion':               '检讨与改进',
    'operation':                '标准化',
    'presentation':             '发表与推广',
    'mentor_names':             '辅导员',
    'participant_names':        '圈员/参与者',
}

PHONE    = input('账号: ').strip()
PASSWORD = input('密码: ').strip()

token = requests.post(f'{BASE}/api/auth/login-with-password',
    json={'phone': PHONE, 'password': PASSWORD}, timeout=10).json()['data']['token']
h = {'Authorization': f'Bearer {token}'}
print('登录成功\n')

rows = []
page = 1
while True:
    data  = requests.get(f'{BASE}/api/admin/registrations/filter',
        params={'competitionId':1,'page':page,'size':50}, headers=h, timeout=30).json()['data']
    items = data.get('content') or data.get('items') or []
    total_pages = data.get('totalPages') or 1

    for item in items:
        if item.get('status') != 'SUBMITTED':
            continue

        reg_id = item['id']
        detail = requests.get(f'{BASE}/api/registrations/{reg_id}',
            headers=h, timeout=15).json().get('data', {})

        reg   = detail.get('registration') or {}
        inst  = detail.get('institution') or {}
        act   = detail.get('activityInfo') or {}
        summ  = detail.get('projectSummary') or {}
        members = detail.get('members') or []

        mentors      = '、'.join(m['name'] for m in members if m.get('role') == 'MENTOR')
        participants = '、'.join(m['name'] for m in members if m.get('role') == 'PARTICIPANT')

        rows.append({
            'registration_id':          reg_id,
            'project_name':             reg.get('projectName', ''),
            'status':                   reg.get('status', ''),
            'submitted_at':             reg.get('submittedAt', ''),
            'group_type':               reg.get('groupType', ''),
            'group_type_label':         GROUP_MAP.get(reg.get('groupType',''), ''),
            'group_code':               reg.get('groupCode', ''),
            'created_at':               reg.get('createdAt', ''),
            'applicant_name':           item.get('applicantName', ''),
            'applicant_phone':          '',   # 接口未返回
            'institution_name':         inst.get('name', ''),
            'institution_level':        inst.get('level', ''),
            'institution_uscc':         inst.get('uscc', ''),
            'institution_region':       inst.get('region', ''),
            'institution_city':         '',   # 接口未返回
            'activity_theme':           act.get('theme', ''),
            'keywords':                 act.get('keywords', ''),
            'avg_work_years':           act.get('avgWorkYears', ''),
            'avg_age':                  act.get('avgAge', ''),
            'cross_department':         '是' if act.get('crossDepartment') else '否',
            'related_to_digital_ai':    '是' if act.get('relatedToDigitalAi') else '否',
            'subject_type_code':        act.get('subjectTypeCode', ''),
            'subject_type_label':       act.get('subjectTypeLabel', ''),
            'subject_type_other':       act.get('subjectTypeOther', ''),
            'method_code':              act.get('methodCode', ''),
            'method_label':             act.get('methodLabel', ''),
            'method_other':             act.get('methodOther', ''),
            'quality_topic_code':       act.get('qualityTopicCode', ''),
            'quality_topic_label':      act.get('qualityTopicLabel', ''),
            'quality_topic_other':      act.get('qualityTopicOther', ''),
            'experience_improve_code':  act.get('experienceImproveCode', ''),
            'experience_improve_label': act.get('experienceImproveLabel', ''),
            'experience_improve_other': act.get('experienceImproveOther', ''),
            'summary_theme':            summ.get('theme', ''),
            'plan':                     summ.get('plan', ''),
            'problem':                  summ.get('problem', ''),
            'action':                   summ.get('action', ''),
            'success':                  summ.get('success', ''),
            'discussion':               summ.get('discussion', ''),
            'operation':                summ.get('operation', ''),
            'presentation':             summ.get('presentation', ''),
            'mentor_names':             mentors,
            'participant_names':        participants,
        })
        print(f'  [{len(rows)}] reg={reg_id}  {item.get("applicantName","")}  {item.get("institutionName","")}')

        if LIMIT and len(rows) >= LIMIT:
            break

    if (LIMIT and len(rows) >= LIMIT) or page >= total_pages:
        break
    page += 1

# 写 Excel
wb = openpyxl.Workbook()
ws = wb.active
ws.title = '报名数据'

headers = list(rows[0].keys())
header_fill = PatternFill('solid', fgColor='4F81BD')
header_font = Font(color='FFFFFF', bold=True)
for ci, col in enumerate(headers, 1):
    cell = ws.cell(row=1, column=ci, value=COL_ZH.get(col, col))
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center')

for ri, row in enumerate(rows, 2):
    for ci, key in enumerate(headers, 1):
        ws.cell(row=ri, column=ci, value=row[key])

# 自动列宽（长文本列限宽）
long_cols = {'plan','problem','action','success','discussion','operation','presentation',
             'mentor_names','participant_names'}
for col in ws.columns:
    key = headers[col[0].column - 1]
    if key in long_cols:
        ws.column_dimensions[col[0].column_letter].width = 50
    else:
        width = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 30)

from datetime import datetime
suffix = f'_前{LIMIT}条' if LIMIT else '_全量'
ts  = datetime.now().strftime('%H%M%S')
out = os.path.join(os.path.dirname(__file__), f'报名数据{suffix}_{ts}.xlsx')
wb.save(out)
print(f'\n导出完成，共 {len(rows)} 条 → {out}')

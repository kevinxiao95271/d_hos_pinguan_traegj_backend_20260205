import openpyxl, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import glob, os

root = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205'
files = [f for f in os.listdir(root) if '项目名称' in f and f.endswith('.xlsx')]
print("找到文件:", files)

for fname in files:
    path = os.path.join(root, fname)
    wb = openpyxl.load_workbook(path, read_only=True)
    for sh in wb.sheetnames:
        ws = wb[sh]
        rows = list(ws.iter_rows(values_only=True))
        print(f"\n=== {fname} / {sh} ===")
        for r in rows:
            if any(c is not None for c in r):
                print(r)

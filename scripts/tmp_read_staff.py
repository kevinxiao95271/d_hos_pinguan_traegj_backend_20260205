import openpyxl, os, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

root = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205'
files = os.listdir(root)
target = [f for f in files if '权限' in f or 'staff' in f.lower()][0]
print(f"文件: {target}")

wb = openpyxl.load_workbook(os.path.join(root, target), read_only=True)
print(f"Sheet: {wb.sheetnames}")
for sh in wb.sheetnames:
    ws = wb[sh]
    rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    print(f"\n=== {sh} ({len(rows)}行) ===")
    for i, row in enumerate(rows[:50]):
        print(f"  {i+1}: {row}")

import openpyxl, re, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

path = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205\4.排程-20260526(1).xlsx'
wb = openpyxl.load_workbook(path, read_only=True)

for shname in wb.sheetnames:
    m = re.match(r'^(6\.[345])', shname)
    if not m:
        continue
    ws = wb[shname]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    # 打印表头和前几行，看列结构
    print(f"\n===== {shname} =====")
    print(f"  表头: {rows[0]}")
    data_rows = [r for r in rows[1:] if r and len(r) >= 3 and isinstance(r[1], (int, float)) and isinstance(r[2], (int, float))]
    for r in data_rows[:5]:
        print(f"  {r}")
    if len(data_rows) > 5:
        print(f"  ... 共 {len(data_rows)} 条")

"""检查汇报书DOCX里是否有活动说明，并打印20260803汇报书的前几个表格"""
import os, sys, re
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

BASE = r'D:\AiCode\cursor\d_hos_pinguan_traegj_backend_20260205\scripts\7error'

def get_folder(reg_id):
    return os.path.join(BASE, [d for d in os.listdir(BASE) if d.startswith(reg_id)][0])

def check_all_docx(reg_id):
    folder = get_folder(reg_id)
    for fn in sorted(os.listdir(folder)):
        if not fn.endswith('.docx') or fn.startswith('~$'):
            continue
        fp = os.path.join(folder, fn)
        print(f'\n  [{reg_id}] {fn}')
        try:
            doc = Document(fp)
            for ti, table in enumerate(doc.tables):
                cells_text = '\n'.join(
                    ' | '.join(c.text.strip() for c in row.cells)
                    for row in table.rows
                )
                if '主题类型' in cells_text and '运用手法' in cells_text:
                    print(f'    *** 表格{ti+1} 包含活动说明 ***')
                    print(cells_text[:1500])
                    return True
            print(f'    (表格数={len(doc.tables)}, 无活动说明)')
        except Exception as e:
            print(f'    失败: {e}')
    return False

for rid in ['20260563', '20260621', '20260803', '20260889']:
    print('='*60)
    check_all_docx(rid)

# 20260803 汇报书前几张表格的内容概况
print('\n\n' + '='*60)
print('20260803 汇报书 前3表格内容')
folder = get_folder('20260803')
fp = os.path.join(folder, '提高ICU危重患者早期康复执行率成果汇报书.docx')
try:
    doc = Document(fp)
    for i, table in enumerate(doc.tables[:3]):
        print(f'\n表格{i+1}:')
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            seen = []
            for c in cells:
                if not seen or c != seen[-1]:
                    seen.append(c)
            line = ' | '.join(seen)
            if line.strip():
                print(f'  {line}')
except Exception as e:
    print(f'读取失败: {e}')

print('\n20260563 汇报书 前3表格内容')
folder = get_folder('20260563')
fp = os.path.join(folder, '"信息化赋能"提高术前去除毛发正确率（汇报书）.docx')
try:
    doc = Document(fp)
    for i, table in enumerate(doc.tables[:3]):
        print(f'\n表格{i+1}:')
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            seen = []
            for c in cells:
                if not seen or c != seen[-1]:
                    seen.append(c)
            line = ' | '.join(seen)
            if line.strip():
                print(f'  {line}')
except Exception as e:
    print(f'读取失败: {e}')

import openpyxl, re, io, sys, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

base = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205'

# 弃赛
wb_w = openpyxl.load_workbook(os.path.join(base, '弃赛项目汇总(1).xlsx'), data_only=True)
withdrawn = set()
for r in wb_w['Sheet1'].iter_rows(values_only=True):
    rid = r[4]
    if isinstance(rid, (int, float)) and rid > 20000000:
        withdrawn.add(int(rid))

# 评分表类型 Excel
wb_f = openpyxl.load_workbook(os.path.join(base, '现场竞赛项目分组及评分表类型 - 发工程师(1).xlsx'), data_only=True)

target_sessions = {
    '综合组-课题达成及QFD专场1改2（三楼开元B厅）',
    '综合组-问题解决型专场1（三楼开元A厅）',
}

for shname in wb_f.sheetnames:
    if not re.match(r'^6\.[345]', shname): continue
    sess_label = re.sub(r'^6\.[345]', '', shname).strip()
    if sess_label not in target_sessions: continue

    ws = wb_f[shname]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    try:
        order_col = next((i for i,h in enumerate(header) if h == '顺序'), None)
        id_col    = next(i for i,h in enumerate(header) if h == '项目编号')
        inst_col  = next(i for i,h in enumerate(header) if h == '机构名称')
        name_col  = next(i for i,h in enumerate(header) if h == '项目名称')
        tool_col  = next((i for i,h in enumerate(header) if h == '运用工具'), None)
        form_col  = next(i for i,h in enumerate(header) if h == '评分表')
    except StopIteration:
        continue

    print(f"\n===== {sess_label} =====")
    print(f"{'顺序':<4} {'项目编号':<10} {'评分表':<8} {'机构名称':<20} 项目名称")
    print("-" * 100)
    count = 0
    for r in rows[1:]:
        if not r or len(r) <= max(id_col, form_col): continue
        rid = r[id_col]
        if not isinstance(rid, (int, float)): continue
        rid = int(rid)
        if rid in withdrawn: continue
        order = r[order_col] if order_col is not None else ''
        inst  = r[inst_col] or ''
        name  = r[name_col] or ''
        fv    = r[form_col] or ''
        count += 1
        print(f"{str(order):<4} {rid:<10} {fv:<8} {str(inst)[:20]:<20} {name}")
    print(f"  共 {count} 项")

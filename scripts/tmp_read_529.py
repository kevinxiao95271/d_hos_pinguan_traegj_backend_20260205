import os, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

root = r'd:\iCode\cursor\d_hos_pinguan_traegj_backend_20260205'
files = os.listdir(root)
targets = [f for f in files if '529' in f or '5.29' in f]
print("目标文件:")
for f in targets:
    print(f"  {f}")

# 用python-docx读取
try:
    import docx
    for fname in targets:
        path = os.path.join(root, fname)
        print(f"\n{'='*60}")
        print(f"文件: {fname}")
        print('='*60)
        doc = docx.Document(path)
        for para in doc.paragraphs:
            if para.text.strip():
                print(para.text)
        for table in doc.tables:
            print(f"\n[表格]")
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    print(' | '.join(cells))
except ImportError:
    print("需要安装 python-docx: pip install python-docx")
